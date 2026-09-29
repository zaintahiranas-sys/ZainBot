"""ZainUzairBot RAG pipeline: load -> clean -> chunk -> embed -> FAISS -> retrieve -> generate."""
import os
import re
import time

import faiss
import numpy as np
from docx import Document
from pypdf import PdfReader

BOT_NAME = "ZainUzairBot"

# Backup Gemini models used automatically when the main model is busy.
FALLBACK_MODELS = ["gemini-3.8-flash", "gemini-3.7-flash", "gemini-3.5-flash", "gemini-3.5-flash-lite"]

SYSTEM_PROMPT = f"""You are {BOT_NAME}, an AI assistant that answers questions about two people:
1. Muhammad Zain, a software engineer from Lahore, Pakistan (documents: Muhammad_Zain_CV and Muhammad_Zain_Resume_v2).
2. Uzair Bin Ahmad, a Software Engineering student and frontend developer from Lahore, Pakistan (document: Uzair_Bin_Ahmad_CV).
You answer questions about their education, experience, projects, skills, teaching and interests.

Rules:
1. Answer ONLY from the CONTEXT provided with each question. Never invent facts.
2. Work out which person the question is about (from the name in the question, or from the chat history for follow-ups). Use the [Source: ...] file name of each context chunk to know whose information it is, and NEVER mix facts of one person into the other. If the question does not say who it is about and it is unclear, briefly ask whether the user means Zain or Uzair.
3. If the answer is not in the context, say: "I don't have that information about <name>." and suggest asking about their projects, skills, education or experience.
4. Speak about each person in the third person. Be friendly, clear and concise.
5. Zain has two CV versions. If they differ (for example dates or job title), prefer the newer resume (Resume v2) and mention the difference only if it matters. This rule applies only to Zain.
6. Never share phone numbers or home address. If asked for contact details, give only the email address if it is in the context.
7. Use the chat history to understand follow-up questions such as "tell me more about it".
8. If asked to compare Zain and Uzair, use only facts from the context for each of them.
9. Reply in the same language the user writes in.
"""


# ---------- 1. Document loading ----------
def _read_pdf(path):
    reader = PdfReader(path)
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _read_docx(path):
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def load_documents(data_dir):
    """Return a list of (filename, raw_text) for every supported file in data_dir."""
    docs = []
    for name in sorted(os.listdir(data_dir)):
        path = os.path.join(data_dir, name)
        ext = name.lower().rsplit(".", 1)[-1]
        if ext == "pdf":
            text = _read_pdf(path)
        elif ext == "docx":
            text = _read_docx(path)
        elif ext in ("txt", "md"):
            with open(path, encoding="utf-8") as f:
                text = f.read()
        else:
            continue
        docs.append((name, text))
    return docs


# ---------- 2. Preprocessing ----------
PHONE_RE = re.compile(r"\(?\+\d{1,3}\)?[\s\-]*\d[\d\s\-]{8,12}\d")


def clean_text(text):
    """Normalize whitespace and redact phone numbers (privacy for a public bot)."""
    text = text.replace("\u00a0", " ").replace("\t", " ")
    text = PHONE_RE.sub("[phone hidden]", text)
    lines = [re.sub(r"\s+", " ", ln).strip() for ln in text.split("\n")]
    return "\n".join(ln for ln in lines if ln)


def chunk_text(text, source, max_chars=650, overlap_lines=2):
    """Group lines into chunks of ~max_chars, overlapping by a couple of lines."""
    lines = text.split("\n")
    chunks, current = [], []
    for line in lines:
        if current and sum(len(x) + 1 for x in current) + len(line) > max_chars:
            chunks.append("\n".join(current))
            current = current[-overlap_lines:]
        current.append(line)
    if current:
        chunks.append("\n".join(current))
    return [f"[Source: {source}]\n{c}" for c in chunks]


# ---------- 3. Knowledge base (embeddings + FAISS) ----------
class KnowledgeBase:
    def __init__(self, api_key, data_dir="data",
                 embed_model="gemini-embedding-001", chat_model="gemini-3.5-flash-lite"):
        from google import genai  # imported here so tests can run without the SDK

        self.client = genai.Client(api_key=api_key)
        self.embed_model = embed_model
        self.chat_model = chat_model
        self.chunks = []
        self.index = None
        self._build(data_dir)

    def _embed(self, texts, task_type):
        from google.genai import types

        vectors = []
        for i in range(0, len(texts), 50):
            batch = texts[i:i + 50]
            for attempt in range(3):
                try:
                    resp = self.client.models.embed_content(
                        model=self.embed_model,
                        contents=batch,
                        config=types.EmbedContentConfig(task_type=task_type),
                    )
                    vectors.extend(e.values for e in resp.embeddings)
                    break
                except Exception:
                    if attempt == 2:
                        raise
                    time.sleep(2 * (attempt + 1))
        arr = np.array(vectors, dtype="float32")
        faiss.normalize_L2(arr)  # cosine similarity via inner product
        return arr

    def _build(self, data_dir):
        for source, raw in load_documents(data_dir):
            self.chunks.extend(chunk_text(clean_text(raw), source))
        vectors = self._embed(self.chunks, "RETRIEVAL_DOCUMENT")
        self.index = faiss.IndexFlatIP(vectors.shape[1])
        self.index.add(vectors)

    def retrieve(self, query, k=5):
        qv = self._embed([query], "RETRIEVAL_QUERY")
        scores, ids = self.index.search(qv, k)
        return [(self.chunks[i], float(s)) for i, s in zip(ids[0], scores[0]) if i >= 0]

    # ---------- 4. Generation with history ----------
    def answer(self, question, history, k=5):
        from google.genai import types

        # Use the previous user question too, so follow-ups retrieve well.
        prev_users = [m["content"] for m in history if m["role"] == "user"]
        search_query = (prev_users[-1] + " " + question) if prev_users else question
        hits = self.retrieve(search_query, k=k)
        context = "\n\n---\n\n".join(c for c, _ in hits)

        contents = []
        for m in history[-8:]:
            role = "user" if m["role"] == "user" else "model"
            contents.append(types.Content(role=role, parts=[types.Part(text=m["content"])]))
        prompt = f"CONTEXT:\n{context}\n\nQUESTION: {question}"
        contents.append(types.Content(role="user", parts=[types.Part(text=prompt)]))

        config = types.GenerateContentConfig(
            system_instruction=SYSTEM_PROMPT, temperature=0.3, max_output_tokens=800
        )
        # Try the main model first; if it is busy (503/429) or unavailable, fall back to the next one.
        models = [self.chat_model] + [m for m in FALLBACK_MODELS if m != self.chat_model]
        resp = None
        for model in models:
            for attempt in range(3):
                try:
                    resp = self.client.models.generate_content(
                        model=model, contents=contents, config=config
                    )
                    break
                except Exception as e:
                    msg = str(e)
                    if any(x in msg for x in ("503", "UNAVAILABLE", "429", "RESOURCE_EXHAUSTED")):
                        time.sleep(2 * (attempt + 1))
                        continue
                    break  # e.g. 404 model not found: move on to the next model
            if resp is not None:
                break
        if resp is None:
            return "The AI model is very busy right now. Please try again in a minute.", hits
        return (resp.text or "Sorry, I could not generate an answer."), hits
